"""
Word document reader module.

Reads participant submissions from ``.docx`` files and returns raw text,
including paragraphs and tables. Falls back gracefully when python-docx is not
installed so the rest of the pipeline can still be imported and tested.
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Union

try:
    import docx  # python-docx
    _HAS_DOCX = True
except Exception:  # pragma: no cover - optional dependency
    _HAS_DOCX = False


class DocumentReadError(RuntimeError):
    """Raised when a submission file cannot be read."""


def _iter_block_text(document) -> List[str]:
    """Extract text from paragraphs and tables in document order-ish.

    python-docx does not expose a single ordered iterator over mixed content,
    so we read all paragraphs, then all tables. For evaluation purposes the
    exact interleaving rarely matters; the full text is what the LLM needs.
    """
    parts: List[str] = []

    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                parts.append(row_text)

    return parts


def read_docx(path: Union[str, Path]) -> str:
    """Read a single ``.docx`` file and return its raw text.

    Raises
    ------
    DocumentReadError
        If python-docx is unavailable or the file cannot be parsed.
    """
    path = Path(path)
    if not path.exists():
        raise DocumentReadError(f"File not found: {path}")
    if path.suffix.lower() != ".docx":
        raise DocumentReadError(
            f"Unsupported file type '{path.suffix}'. Only .docx is supported "
            "(convert .doc/.pdf first)."
        )
    if not _HAS_DOCX:
        raise DocumentReadError(
            "python-docx is not installed. Run `pip install python-docx`."
        )

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentReadError(f"Could not open '{path}': {exc}") from exc

    return "\n".join(_iter_block_text(document))


def discover_submission_files(folder: Union[str, Path]) -> List[Path]:
    """Return sorted ``.docx`` files in a folder (non-recursive by default)."""
    folder = Path(folder)
    if not folder.is_dir():
        raise DocumentReadError(f"Not a directory: {folder}")
    files = sorted(
        p for p in folder.glob("*.docx")
        if not p.name.startswith("~$")  # skip Word lock/temp files
    )
    return files
