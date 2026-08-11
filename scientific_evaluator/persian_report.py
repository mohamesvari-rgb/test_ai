"""
Persian (RTL) Word renderer.

Renders the shared report blocks (see report_builder) into a right-to-left Word
document. Persian text is tagged with the requested font (default "B Nazanin",
as the complex-script font) and Latin text/numbers with "Times New Roman", so
the .docx honours the requested fonts exactly on any machine that has them.

A thin ``PersianReportGenerator`` builds the blocks and renders them; charts are
generated on demand.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import List, Optional, Union

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import charts as charts_mod
from .config import SETTINGS
from .models import EvaluationResult
from .report_builder import (
    Block, Bullets, Heading, Image, Para, Table, Title, build_report_blocks,
)

_HEADING_COLOR = RGBColor(0x1F, 0x4E, 0x79)
_HEADER_SHADE = "D9E2F3"


# ---------------------------------------------------------------------------
# Low-level RTL / font helpers
# ---------------------------------------------------------------------------
def _style_run(run, size: int, bold: bool = False, color: Optional[RGBColor] = None):
    """Apply Persian(cs)+Latin fonts and mark the run right-to-left."""
    persian = SETTINGS.persian_font
    latin = SETTINGS.latin_font
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color

    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:cs"), persian)      # complex-script (Persian) font

    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rPr.append(rtl)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)
    if bold:
        rPr.append(OxmlElement("w:bCs"))


def _rtl_paragraph(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))
    if align is not None:
        p.alignment = align


def _add_par(doc, text, size=None, bold=False, color=None,
             align=WD_ALIGN_PARAGRAPH.RIGHT):
    size = size or SETTINGS.base_font_size
    p = doc.add_paragraph()
    _rtl_paragraph(p, align)
    if text:
        _style_run(p.add_run(text), size, bold, color)
    return p


def _shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell(cell, text, header=False, size=None):
    size = size or (SETTINGS.base_font_size - 1)
    cell.text = ""
    p = cell.paragraphs[0]
    _rtl_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
    for idx, line in enumerate(str(text).split("\n")):
        if idx > 0:
            p = cell.add_paragraph()
            _rtl_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
        _style_run(p.add_run(line), size, header,
                   _HEADING_COLOR if header else None)
    if header:
        _shade_cell(cell, _HEADER_SHADE)


# ---------------------------------------------------------------------------
# Block renderer
# ---------------------------------------------------------------------------
def _render_block(doc, block: Block):
    if isinstance(block, Title):
        _add_par(doc, block.text, size=22, bold=True, color=_HEADING_COLOR,
                 align=WD_ALIGN_PARAGRAPH.CENTER)
        for sub in block.subtitles:
            _add_par(doc, sub, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _add_par(doc, "")

    elif isinstance(block, Heading):
        h = doc.add_heading("", level=block.level)
        _rtl_paragraph(h)
        size = SETTINGS.heading_font_size - (block.level - 1)
        _style_run(h.add_run(block.text), size, bold=True, color=_HEADING_COLOR)

    elif isinstance(block, Para):
        _add_par(doc, block.text, bold=block.bold)

    elif isinstance(block, Bullets):
        for it in block.items:
            if str(it).strip():
                _add_par(doc, f"• {it}")

    elif isinstance(block, Table):
        table = doc.add_table(rows=1, cols=len(block.headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.RIGHT
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))  # RTL column order
        for i, h in enumerate(block.headers):
            _set_cell(table.rows[0].cells[i], h, header=True)
        for row in block.rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                if i < len(cells):
                    _set_cell(cells[i], val)
        if block.widths:
            total = Inches(6.6)
            for r in table.rows:
                for i, w in enumerate(block.widths):
                    if i < len(r.cells):
                        r.cells[i].width = Inches(6.6 * w)

    elif isinstance(block, Image):
        if Path(block.path).exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(block.path, width=Inches(6.2))
            if block.caption:
                _add_par(doc, block.caption, size=10,
                         align=WD_ALIGN_PARAGRAPH.CENTER)


def render_blocks_to_docx(blocks: List[Block], out_path: Union[str, Path]) -> Path:
    """Render report blocks to a Persian RTL .docx file."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = SETTINGS.latin_font
    normal.font.size = Pt(SETTINGS.base_font_size)

    for block in blocks:
        _render_block(doc, block)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# Convenience generator (builds blocks + charts, then renders docx)
# ---------------------------------------------------------------------------
class PersianReportGenerator:
    """Builds the report blocks (with charts) and renders a Persian .docx."""

    def generate(
        self,
        result: EvaluationResult,
        out_path: Union[str, Path],
        chart_dir: Optional[Path] = None,
    ) -> Path:
        tmp_ctx = None
        chart_paths: List[Path] = []
        if charts_mod.charts_available():
            if chart_dir is None:
                tmp_ctx = tempfile.TemporaryDirectory()
                chart_dir = Path(tmp_ctx.name)
            chart_paths = charts_mod.generate_charts(result, chart_dir)
        try:
            blocks = build_report_blocks(result, chart_paths)
            return render_blocks_to_docx(blocks, out_path)
        finally:
            if tmp_ctx is not None:
                tmp_ctx.cleanup()


# ---------------------------------------------------------------------------
# Optional: higher-fidelity PDF via LibreOffice (honours installed B Nazanin).
# The default PDF path uses weasyprint (see pdf_report.py); this is a fallback
# for environments where LibreOffice is available and the exact fonts installed.
# ---------------------------------------------------------------------------
def convert_docx_to_pdf_libreoffice(
    docx_path: Union[str, Path], pdf_path: Optional[Union[str, Path]] = None
) -> Path:
    """Convert a .docx to .pdf using LibreOffice, if available and working."""
    docx_path = Path(docx_path)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) not found.")
    out_dir = Path(pdf_path).parent if pdf_path else docx_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [soffice, f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
            check=True, capture_output=True, timeout=180,
        )
    produced = out_dir / (docx_path.stem + ".pdf")
    if not produced.exists():
        raise RuntimeError(f"LibreOffice did not produce {produced}")
    if pdf_path and Path(pdf_path) != produced:
        shutil.move(str(produced), str(pdf_path))
        produced = Path(pdf_path)
    return produced
