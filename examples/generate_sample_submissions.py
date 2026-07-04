#!/usr/bin/env python3
"""
Generate sample participant submissions as .docx files for testing/demo.

Run:
    python examples/generate_sample_submissions.py
Creates:
    examples/submissions/participant_A.docx  (strong)
    examples/submissions/participant_B.docx  (medium)
    examples/submissions/participant_C.docx  (weak)

These let you run the full pipeline end-to-end (with --mock or a real API key)
without needing real participant files.
"""

from __future__ import annotations

from pathlib import Path

try:
    import docx
except ImportError:
    raise SystemExit("python-docx is required: pip install python-docx")


SUBMISSIONS = {
    "participant_A": {
        "title": "Sensor-Fusion Water Stress Detection with Self-Calibrating Edge AI",
        "paragraphs": [
            "We propose a hybrid system that fuses free Sentinel-2 multispectral "
            "imagery (10-20 m, 5-day revisit) with a sparse network of low-cost "
            "soil-moisture and canopy-temperature sensors (one node per ~5 farms).",
            "Core idea: the ground nodes continuously self-calibrate the satellite "
            "model using a lightweight Gaussian-process residual correction, so no "
            "expert calibration is needed per farm. This addresses the scaling "
            "constraint directly.",
            "Water stress is predicted from a physically-grounded Crop Water Stress "
            "Index (CWSI) augmented by an ML model trained on canopy temperature "
            "and NDRE/NDWI indices. We target 3-5 day lead time before visible "
            "wilting, validated against a held-out season of ground truth.",
            "Edge nodes run on solar + battery and cache alerts locally, syncing "
            "over intermittent connectivity via LoRaWAN to a gateway. Alerts are "
            "delivered as simple SMS with an irrigation recommendation.",
            "Validation plan: 2 growing seasons, 30 plots, compare predicted stress "
            "onset to pressure-chamber leaf water potential; target >85% precision "
            "at 4-day lead time. Cost: ~$40 per farm amortized (shared nodes).",
            "Risks: LoRa range in hilly terrain, sensor drift; mitigated by "
            "redundant nodes and periodic self-calibration.",
        ],
    },
    "participant_B": {
        "title": "Satellite NDVI Thresholding for Irrigation Alerts",
        "paragraphs": [
            "Our approach uses Sentinel-2 NDVI time series to detect drops that "
            "indicate water stress. When NDVI falls below a threshold, we send an "
            "alert to the farmer.",
            "We will build a mobile app that shows the NDVI map and a red/green "
            "status per field. The model is a threshold tuned on last year's data.",
            "This is low cost because it only uses free satellite data and does not "
            "require ground sensors. It should work for many farms.",
            "We expect reasonable accuracy but have not yet quantified lead time. "
            "Connectivity is handled by the mobile app caching maps.",
        ],
    },
    "participant_C": {
        "title": "AI-Powered Smart Farming Platform",
        "paragraphs": [
            "We will build an AI platform that uses big data and deep learning to "
            "revolutionize farming and detect all crop problems including water "
            "stress with very high accuracy.",
            "The system uses drones, satellites, IoT and blockchain to give farmers "
            "a complete solution. Our neural network is state of the art.",
            "It will be cheap and scale to millions of farms. Farmers will love the "
            "dashboard. This will transform agriculture.",
        ],
    },
}


def main() -> None:
    out_dir = Path(__file__).parent / "submissions"
    out_dir.mkdir(parents=True, exist_ok=True)
    for pid, content in SUBMISSIONS.items():
        doc = docx.Document()
        doc.add_heading(content["title"], level=1)
        for para in content["paragraphs"]:
            doc.add_paragraph(para)
        path = out_dir / f"{pid}.docx"
        doc.save(str(path))
        print(f"wrote {path}")


if __name__ == "__main__":
    main()
