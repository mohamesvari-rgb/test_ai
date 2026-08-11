"""
Unit tests for the scoring engine, rubric integrity, and end-to-end pipeline
(using the offline MockLLMClient so tests need no API key or network).

Run:  python -m pytest tests/ -v      (or)   python tests/test_scoring.py
"""

from __future__ import annotations

import sys
from pathlib import Path

# Make the package importable when running the file directly.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scientific_evaluator.config import RUBRIC
from scientific_evaluator.models import Challenge, Submission
from scientific_evaluator.scoring_engine import (
    build_criterion_scores,
    compute_total,
    integrity_penalty,
)
from scientific_evaluator.llm_client import MockLLMClient, extract_json
from scientific_evaluator.agent import ScientificEvaluationAgent


def test_rubric_weights_sum_to_100():
    assert abs(sum(c.max_points for c in RUBRIC) - 100.0) < 1e-6


def test_perfect_scores_give_100():
    raw = {c.key: {"raw_score": 10, "justification": "x"} for c in RUBRIC}
    scores = build_criterion_scores(raw)
    assert abs(compute_total(scores) - 100.0) < 1e-6


def test_zero_scores_give_0():
    raw = {c.key: {"raw_score": 0, "justification": ""} for c in RUBRIC}
    scores = build_criterion_scores(raw)
    assert compute_total(scores) == 0.0


def test_out_of_range_scores_are_clamped():
    raw = {c.key: {"raw_score": 999, "justification": ""} for c in RUBRIC}
    scores = build_criterion_scores(raw)
    # Clamped to 10 each -> 100 total.
    assert abs(compute_total(scores) - 100.0) < 1e-6


def test_missing_criterion_defaults_to_zero():
    raw = {RUBRIC[0].key: {"raw_score": 10, "justification": ""}}
    scores = build_criterion_scores(raw)
    assert len(scores) == len(RUBRIC)
    # Only the first criterion contributes.
    assert abs(compute_total(scores) - RUBRIC[0].max_points) < 1e-6


def test_integrity_penalty_is_bounded():
    assert integrity_penalty([]) == 0.0
    assert integrity_penalty(["a"]) == 5.0
    assert integrity_penalty(["a", "b", "c", "d", "e", "f"]) == 25.0  # capped


def test_extract_json_handles_fences():
    text = 'Here you go:\n```json\n{"a": 1, "b": [2,3]}\n```\nDone.'
    data = extract_json(text)
    assert data == {"a": 1, "b": [2, 3]}


def test_weighted_points_math():
    raw = {c.key: {"raw_score": 5, "justification": ""} for c in RUBRIC}
    scores = build_criterion_scores(raw)
    # 5/10 of every criterion -> 50.0
    assert abs(compute_total(scores) - 50.0) < 1e-6


def _demo_result(n=3):
    challenge = Challenge(
        title="Test challenge",
        field="testing",
        background="A background.",
        objectives=["obj1", "obj2"],
    )
    subs = [
        Submission(
            participant_id=f"P{i}",
            source_path=f"P{i}.docx",
            raw_text=f"This is submission {i} with some scientific content.",
            title=f"Title {i}",
            word_count=8,
        )
        for i in range(n)
    ]
    agent = ScientificEvaluationAgent(challenge, MockLLMClient())
    return agent, agent.run(subs), subs


def test_end_to_end_pipeline_with_mock():
    _agent, result, subs = _demo_result(3)

    # Every submission is evaluated and scored in range.
    assert len(result.evaluations) == 3
    for e in result.evaluations:
        assert 0.0 <= e.total_score <= 100.0
        assert len(e.criterion_scores) == len(RUBRIC)

    # Ranking covers every participant exactly once.
    assert sorted(result.comparison.ranking) == sorted(e.participant_id for e in subs)

    # Synthesis and recommendation are produced.
    assert result.synthesized_solution
    assert result.final_recommendation


def test_idea_improvements_are_generated():
    _agent, result, _subs = _demo_result(2)
    for e in result.evaluations:
        assert e.idea_improvements, "each submission should get idea improvements"
        for imp in e.idea_improvements:
            assert imp.idea
            assert imp.improvements  # at least one improvement solution


def test_mock_ranking_is_ordered_by_score():
    """The offline demo's ranking must not contradict its own score table."""
    _agent, result, _subs = _demo_result(4)
    ranked = result.ranked_evaluations()
    scores = [e.total_score for e in ranked]
    assert scores == sorted(scores, reverse=True), (
        f"ranking is not strongest-first: {[(e.participant_id, e.total_score) for e in ranked]}"
    )


def test_persian_participant_ids_survive_ranking():
    """Regression: non-ASCII ids must not be \\uXXXX-escaped into the prompt.

    When they were, the model echoed back ids that no longer matched the real
    ones, they were dropped from the ranking, and the report order became
    incoherent (rank 1 scoring lower than rank 2).
    """
    from scientific_evaluator.prompts import build_comparison_prompt

    challenge = Challenge("t", "f", "b", ["o1"])
    subs = [
        Submission("علی_رضایی", "a.docx", "متن علمی با روش و اعتبارسنجی.", "پاسخ", 6),
        Submission("Sara_Team", "b.docx", "A scientific answer.", "Ans", 4),
        Submission("محمد_احمدی", "c.docx", "متن دیگر.", "پاسخ ۲", 3),
    ]
    agent = ScientificEvaluationAgent(challenge, MockLLMClient())
    result = agent.run(subs)

    # Ids reach the prompt as real characters.
    _sys, user = build_comparison_prompt(challenge, result.evaluations)
    assert "علی_رضایی" in user and "\\u0639" not in user

    # Every participant appears exactly once, ordered strongest-first.
    assert sorted(result.comparison.ranking) == sorted(s.participant_id for s in subs)
    scores = [e.total_score for e in result.ranked_evaluations()]
    assert scores == sorted(scores, reverse=True), scores


def test_comparison_has_salient_and_distinctive_points():
    _agent, result, subs = _demo_result(3)
    comp = result.comparison
    for s in subs:
        assert s.participant_id in comp.salient_points
        assert s.participant_id in comp.distinctive_points


def test_persian_report_export_docx(tmp_path=None):
    import tempfile
    from scientific_evaluator.persian_report import PersianReportGenerator
    _agent, result, _subs = _demo_result(3)
    out_dir = Path(tempfile.mkdtemp())
    out = PersianReportGenerator().generate(result, out_dir / "fa_report.docx")
    assert out.exists() and out.stat().st_size > 0

    # Verify the document actually contains Persian section headings + tables.
    import docx
    doc = docx.Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "گزارش ارزیابی علمی" in text          # cover title
    assert "راه‌حل جامع" in text                  # integrated solution section
    assert len(doc.tables) >= 4                    # rubric+scores+innovation+more


def test_word_still_delivered_when_pdf_engine_missing():
    """Regression: a missing GTK/WeasyPrint stack must not lose the whole run.

    Reproduces the Windows failure
    ``OSError: cannot load library 'libgobject-2.0-0'`` and asserts that
    fmt='both' still returns the Word report plus an actionable warning.
    """
    import tempfile
    from scientific_evaluator import agent as agent_mod

    _agent, result, _subs = _demo_result(2)
    messages = []
    _agent._progress = messages.append

    original = agent_mod.render_blocks_to_pdf

    def boom(_blocks, _path):
        raise OSError("cannot load library 'libgobject-2.0-0': error 0x7e")

    agent_mod.render_blocks_to_pdf = boom
    try:
        out = Path(tempfile.mkdtemp()) / "report"
        written = _agent.export_report(result, out, fmt="both")
    finally:
        agent_mod.render_blocks_to_pdf = original

    # The Word file survives and is returned.
    assert len(written) == 1 and written[0].suffix == ".docx"
    assert written[0].exists() and written[0].stat().st_size > 0

    # The warning explains the fix rather than leaking a bare OSError.
    warning = "\n".join(messages)
    assert "WARNING" in warning
    assert "Word only" in warning


def test_pdf_error_message_is_actionable():
    from scientific_evaluator.pdf_report import pdf_engine_error_message

    msg = pdf_engine_error_message(
        OSError("cannot load library 'libgobject-2.0-0': error 0x7e")
    )
    assert "GTK" in msg or "gtk" in msg          # names what is missing
    assert "Word only" in msg                    # gives an immediate workaround


def test_persian_report_export_pdf():
    """PDF export via weasyprint (skipped if weasyprint is unavailable)."""
    try:
        import weasyprint  # noqa: F401
    except Exception:
        print("SKIP test_persian_report_export_pdf (weasyprint not installed)")
        return
    import tempfile
    from scientific_evaluator import charts as charts_mod
    from scientific_evaluator.pdf_report import render_blocks_to_pdf
    from scientific_evaluator.report_builder import build_report_blocks

    _agent, result, _subs = _demo_result(3)
    with tempfile.TemporaryDirectory() as d:
        charts = (charts_mod.generate_charts(result, d)
                  if charts_mod.charts_available() else [])
        blocks = build_report_blocks(result, charts)
        out = render_blocks_to_pdf(blocks, Path(d) / "fa_report.pdf")
        data = out.read_bytes()
        assert data[:4] == b"%PDF"          # valid PDF header
        assert len(data) > 3000             # non-trivial content


if __name__ == "__main__":
    # Minimal runner without pytest.
    import traceback

    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"PASS  {t.__name__}")
        except Exception:  # noqa: BLE001
            failed += 1
            print(f"FAIL  {t.__name__}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    raise SystemExit(1 if failed else 0)
