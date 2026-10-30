"""
Tests for the web GUI (Flask app).

Uses the offline mock backend so no API key or network is needed. Exercises the
full browser flow: load the page, POST a challenge + .docx uploads, poll until
the job finishes, and download the generated report.

Run:  python tests/test_webapp.py     (or)  python -m pytest tests/ -v
"""

from __future__ import annotations

import io
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from webapp.app import create_app, safe_filename  # noqa: E402


def _docx_bytes(text: str = "A scientific answer with a method and validation.") -> bytes:
    """Build a minimal in-memory .docx."""
    import docx
    doc = docx.Document()
    doc.add_heading("Proposed Solution", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _client(tmp_root):
    app = create_app(work_root=tmp_root)
    app.config["TESTING"] = True
    return app.test_client()


def test_safe_filename_preserves_persian():
    # Persian names must survive (the stem becomes the participant id).
    assert safe_filename("علی_رضایی.docx") == "علی_رضایی.docx"
    # Path traversal is neutralized.
    assert "/" not in safe_filename("../../etc/passwd.docx")
    assert safe_filename("../../etc/passwd.docx") == "passwd.docx"
    # Extension is enforced.
    assert safe_filename("answer").endswith(".docx")


def test_index_page_renders():
    import tempfile
    c = _client(tempfile.mkdtemp())
    r = c.get("/")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "Scientific Evaluation Agent" in body
    assert 'id="form"' in body


def test_engine_endpoint_reports_both_modes():
    import tempfile
    c = _client(tempfile.mkdtemp())
    data = c.get("/api/engine").get_json()
    assert data["mode"] in ("online", "offline", "demo")
    assert set(data["online"]) == {"provider", "has_key"}
    assert set(data["offline"]) == {"running", "model_installed"}


def test_offline_engine_needs_no_api_key():
    """Choosing offline starts a job without a key; a stopped server is reported."""
    import os
    import tempfile
    saved = os.environ.get("OLLAMA_BASE_URL")
    os.environ["OLLAMA_BASE_URL"] = "http://127.0.0.1:9"
    try:
        c = _client(tempfile.mkdtemp())
        assert c.get("/api/engine").get_json()["offline"]["running"] is False
        r = c.post("/api/evaluate", data={
            "title": "t", "field": "f", "background": "b", "objectives": "o1",
            "provider": "ollama",
            "submissions": (io.BytesIO(_docx_bytes()), "answer.docx"),
        }, content_type="multipart/form-data")
        assert r.status_code == 200, r.get_data(as_text=True)
        job_id = r.get_json()["job_id"]
        for _ in range(60):
            s = c.get(f"/api/status/{job_id}").get_json()
            if s["status"] in ("done", "error"):
                break
            time.sleep(0.5)
        assert s["status"] == "error" and "not running" in s["error"], s["error"]
    finally:
        if saved is None:
            os.environ.pop("OLLAMA_BASE_URL", None)
        else:
            os.environ["OLLAMA_BASE_URL"] = saved


def test_missing_fields_rejected():
    import tempfile
    c = _client(tempfile.mkdtemp())
    r = c.post("/api/evaluate", data={"title": "only a title", "mock": "on"},
               content_type="multipart/form-data")
    assert r.status_code == 400
    assert "required" in r.get_json()["error"].lower()


def test_non_docx_rejected():
    import tempfile
    c = _client(tempfile.mkdtemp())
    r = c.post("/api/evaluate", data={
        "title": "t", "field": "f", "background": "b", "objectives": "o1",
        "mock": "on",
        "submissions": (io.BytesIO(b"not a docx"), "answer.pdf"),
    }, content_type="multipart/form-data")
    assert r.status_code == 400
    assert ".docx" in r.get_json()["error"]


def test_full_gui_flow_with_mock():
    """End-to-end: submit two answers, poll to completion, download the report."""
    import tempfile
    c = _client(tempfile.mkdtemp())

    r = c.post("/api/evaluate", data={
        "title": "Detect crop water stress",
        "field": "remote sensing",
        "background": "Smallholder farms lose yield to late-detected water stress.",
        "objectives": "Detect stress early\nKeep cost low",
        "constraints": "Intermittent connectivity",
        "expected_outputs": "System architecture",
        "evaluation_notes": "Reward validated, low-cost approaches.",
        "language": "Persian",
        "format": "both",
        "mock": "on",
        "submissions": [
            (io.BytesIO(_docx_bytes("Sensor fusion with self-calibration.")), "علی_رضایی.docx"),
            (io.BytesIO(_docx_bytes("NDVI thresholding from satellite data.")), "participant_B.docx"),
        ],
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    job_id = r.get_json()["job_id"]

    # Poll until the background job finishes.
    status = None
    for _ in range(120):                      # up to ~60s
        s = c.get(f"/api/status/{job_id}").get_json()
        status = s["status"]
        if status in ("done", "error"):
            break
        time.sleep(0.5)
    assert status == "done", f"job did not finish: {s.get('error')} | {s.get('messages')}"

    # Both report formats plus the raw JSON should be offered.
    names = s["files"]
    assert any(n.endswith(".docx") for n in names), names
    assert any(n.endswith(".pdf") for n in names), names
    assert any(n.endswith(".json") for n in names), names

    # Ranking covers both participants (Persian filename preserved).
    ids = {row["participant_id"] for row in s["ranking"]}
    assert ids == {"علی_رضایی", "participant_B"}, ids

    # Downloads work and return real files.
    pdf_name = next(n for n in names if n.endswith(".pdf"))
    d = c.get(f"/api/download/{job_id}/{pdf_name}")
    assert d.status_code == 200
    assert d.get_data()[:4] == b"%PDF"

    # Unknown files are not served (path traversal / arbitrary read).
    assert c.get(f"/api/download/{job_id}/secrets.txt").status_code == 404

    # Cleanup removes the job.
    assert c.post(f"/api/cleanup/{job_id}").status_code == 200
    assert c.get(f"/api/status/{job_id}").status_code == 404


def _run_job(c, language, fmt):
    r = c.post("/api/evaluate", data={
        "title": "t", "field": "f", "background": "b", "objectives": "o1",
        "language": language, "format": fmt, "mock": "on",
        "submissions": [(io.BytesIO(_docx_bytes()), "answer.docx")],
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    job_id = r.get_json()["job_id"]
    for _ in range(120):
        s = c.get(f"/api/status/{job_id}").get_json()
        if s["status"] in ("done", "error"):
            return job_id, s
        time.sleep(0.5)
    raise AssertionError("job did not finish")


def test_english_pdf_only_report():
    """Regression: English + PDF stopped the run after the whole evaluation."""
    import tempfile
    c = _client(tempfile.mkdtemp())
    job_id, s = _run_job(c, "English", "pdf")
    assert s["status"] == "done", s.get("error")
    assert sorted(s["files"]) == ["report.pdf", "result.json"], s["files"]
    assert c.get(f"/api/download/{job_id}/report.pdf").get_data()[:4] == b"%PDF"


def test_failed_report_keeps_the_results():
    import tempfile
    from scientific_evaluator.agent import ScientificEvaluationAgent

    def boom(*_args, **_kwargs):
        raise RuntimeError("no renderer")

    original = ScientificEvaluationAgent.export_report
    ScientificEvaluationAgent.export_report = boom
    try:
        _job_id, s = _run_job(_client(tempfile.mkdtemp()), "Persian", "both")
    finally:
        ScientificEvaluationAgent.export_report = original
    assert s["status"] == "done", s.get("error")
    assert s["files"] == ["result.json"]
    assert s["ranking"] and s["summary"]
    assert any(m.startswith("WARNING: the report file") for m in s["messages"])


if __name__ == "__main__":
    import traceback
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"PASS  {t.__name__}")
        except Exception:
            failed += 1
            print(f"FAIL  {t.__name__}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    raise SystemExit(1 if failed else 0)
